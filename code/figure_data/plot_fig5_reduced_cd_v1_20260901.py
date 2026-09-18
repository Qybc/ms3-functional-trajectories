#!/usr/bin/env python3
"""Generate reduced Fig. 5 panels c-d from the supplied raw workbooks.

Panel c retains tensile adhesion and 180-degree peel only, and each chart
contains Fr-PA-10 and PTPU-10 only. Panel d retains the 2,000-cycle electrical
stability readout with raw waveform insets. The SVG contains no clip paths.
"""

from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
PACKAGE_ROOT = HERE.parent.parent
ROOT = HERE / "outputs"
RAW = PACKAGE_ROOT / "data/fig5_cardiac_interface/raw"
OUTPUT = ROOT / "fig5_reduced_cd_v1_20260901.svg"

W, H = 1800, 900
INK = "#111111"
FR = "#315D70"
PTPU = "#F2A65A"
LOWER = "#E66F51"
PALE = "#DCE8EC"


def numeric_pairs(ws, x_col, y_col, min_row=2):
    pairs = []
    for row in ws.iter_rows(min_row=min_row, values_only=True):
        x = row[x_col - 1] if x_col <= len(row) else None
        y = row[y_col - 1] if y_col <= len(row) else None
        if isinstance(x, (int, float)) and isinstance(y, (int, float)):
            pairs.append((float(x), float(y)))
    return pairs


def path_d(points):
    if not points:
        return ""
    return "M " + " L ".join(f"{x:.2f},{y:.2f}" for x, y in points)


def text(x, y, value, size=22, anchor="middle", weight=400, rotate=None, color=INK):
    transform = f' transform="rotate({rotate} {x:.2f} {y:.2f})"' if rotate is not None else ""
    return (
        f'<text x="{x:.2f}" y="{y:.2f}" font-size="{size}" font-weight="{weight}" '
        f'text-anchor="{anchor}" fill="{color}"{transform}>{escape(str(value))}</text>'
    )


def line(x1, y1, x2, y2, stroke=INK, width=3.0, opacity=1.0, dash=None):
    dash_attr = f' stroke-dasharray="{dash}"' if dash else ""
    return (
        f'<line x1="{x1:.2f}" y1="{y1:.2f}" x2="{x2:.2f}" y2="{y2:.2f}" '
        f'stroke="{stroke}" stroke-width="{width}" opacity="{opacity}"{dash_attr}/>'
    )


def rect(x, y, width, height, fill="white", stroke=INK, stroke_width=4.0):
    return (
        f'<rect x="{x:.2f}" y="{y:.2f}" width="{width:.2f}" height="{height:.2f}" '
        f'fill="{fill}" stroke="{stroke}" stroke-width="{stroke_width}"/>'
    )


def map_point(x, y, box, xlim, ylim):
    bx, by, bw, bh = box
    px = bx + (x - xlim[0]) / (xlim[1] - xlim[0]) * bw
    py = by + bh - (y - ylim[0]) / (ylim[1] - ylim[0]) * bh
    return px, py


def draw_axes(svg, box, xlim, ylim, xticks, yticks, xlabel, ylabel, tick_size=19):
    x, y, w, h = box
    for value in yticks[1:-1]:
        _, py = map_point(xlim[0], value, box, xlim, ylim)
        svg.append(line(x, py, x + w, py, width=1.7, opacity=0.82, dash="10 9"))
    svg.append(rect(x, y, w, h, fill="white", stroke=INK, stroke_width=4.2))
    for value in xticks:
        px, _ = map_point(value, ylim[0], box, xlim, ylim)
        svg.append(line(px, y + h, px, y + h + 9, width=3.2))
        svg.append(text(px, y + h + 31, f"{value:g}", size=tick_size))
    for value in yticks:
        _, py = map_point(xlim[0], value, box, xlim, ylim)
        svg.append(line(x - 9, py, x, py, width=3.2))
        svg.append(text(x - 15, py + 6, f"{value:g}", size=tick_size, anchor="end"))
    svg.append(text(x + w / 2, y + h + 68, xlabel, size=26, weight=600))
    svg.append(text(x - 66, y + h / 2, ylabel, size=25, weight=600, rotate=-90))


def draw_curve(svg, values, box, xlim, ylim, color):
    points = [
        map_point(x, y, box, xlim, ylim)
        for x, y in values
        if xlim[0] <= x <= xlim[1] and ylim[0] <= y <= ylim[1]
    ]
    svg.append(
        f'<path d="{path_d(points)}" fill="none" stroke="{color}" stroke-width="5.0" '
        f'stroke-linejoin="round" stroke-linecap="round"/>'
    )


def draw_legend(svg, x, y):
    svg.append(line(x, y, x + 55, y, stroke=FR, width=5.0))
    svg.append(text(x + 72, y + 7, "Fr-PA-10", size=20, anchor="start", weight=500))
    svg.append(line(x, y + 38, x + 55, y + 38, stroke=PTPU, width=5.0))
    svg.append(text(x + 72, y + 45, "PTPU-10", size=20, anchor="start", weight=500))


def percentile(values, fraction):
    ordered = sorted(values)
    position = fraction * (len(ordered) - 1)
    lo = int(position)
    hi = min(lo + 1, len(ordered) - 1)
    weight = position - lo
    return ordered[lo] * (1 - weight) + ordered[hi] * weight


def cycle_coordinate(t, last_time, total=2000):
    return t / last_time * total


def windowed_envelope(raw, total=2000, window_width=10):
    last_time = raw[-1][0]
    windows = [[] for _ in range(total // window_width)]
    for time_value, resistance in raw:
        coordinate = cycle_coordinate(time_value, last_time, total)
        index = int(coordinate // window_width)
        if 0 <= index < len(windows):
            windows[index].append(resistance)
    upper, lower = [], []
    for index, values in enumerate(windows):
        if len(values) < 20:
            continue
        x = index * window_width + window_width / 2
        upper.append((x, percentile(values, 0.95)))
        lower.append((x, percentile(values, 0.05)))
    return upper, lower


def draw_inset(svg, raw, center_cycle, box):
    x, y, w, h = box
    lo, hi = center_cycle - 5, center_cycle + 5
    last_time = raw[-1][0]
    values = [
        (cycle_coordinate(t, last_time), resistance)
        for t, resistance in raw
        if lo <= cycle_coordinate(t, last_time) <= hi
    ]
    xlim, ylim = (lo, hi), (1.4, 2.0)
    for value in [1.7]:
        _, py = map_point(lo, value, box, xlim, ylim)
        svg.append(line(x, py, x + w, py, width=1.35, opacity=0.75, dash="8 7"))
    svg.append(rect(x, y, w, h, fill="white", stroke=INK, stroke_width=3.5))
    for value in [lo, center_cycle, hi]:
        px, _ = map_point(value, ylim[0], box, xlim, ylim)
        svg.append(line(px, y + h, px, y + h + 7, width=2.5))
        svg.append(text(px, y + h + 24, f"{value:g}", size=15))
    for value in [1.4, 1.7, 2.0]:
        _, py = map_point(lo, value, box, xlim, ylim)
        svg.append(line(x - 7, py, x, py, width=2.5))
        svg.append(text(x - 12, py + 5, f"{value:g}", size=15, anchor="end"))
    points = [map_point(cx, r, box, xlim, ylim) for cx, r in values]
    svg.append(
        f'<path d="{path_d(points)}" fill="none" stroke="{FR}" stroke-width="3.4" '
        f'stroke-linejoin="round" stroke-linecap="round"/>'
    )


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    if OUTPUT.exists():
        raise SystemExit(f"Refusing to overwrite existing output: {OUTPUT}")

    adhesion = RAW
    tensile_ws = load_workbook(adhesion / "tensile_adhesion_raw.xlsx", read_only=True, data_only=True).active
    peel_ws = load_workbook(adhesion / "peel_180_degree_raw.xlsx", read_only=True, data_only=True).active
    tensile_fr = numeric_pairs(tensile_ws, 4, 5, min_row=2)
    tensile_ptpu = numeric_pairs(tensile_ws, 10, 11, min_row=2)
    peel_fr = numeric_pairs(peel_ws, 4, 5, min_row=2)
    peel_ptpu = numeric_pairs(peel_ws, 10, 11, min_row=2)

    cycle_ws = load_workbook(
        RAW / "cyclic_resistance_2000cycles_50pct_raw.xlsx",
        read_only=True, data_only=True,
    ).active
    cycle_raw = numeric_pairs(cycle_ws, 1, 2, min_row=2)
    upper, lower = windowed_envelope(cycle_raw)

    svg = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">',
        '<defs><style><![CDATA[',
        "text { font-family: 'Avenir Next', Avenir, Arial, sans-serif; fill: #111111; }",
        ']]></style></defs>',
        f'<rect x="0" y="0" width="{W}" height="{H}" fill="white"/>',
        text(25, 38, "c", size=36, anchor="start", weight=700),
    ]

    tensile_box = (95, 55, 690, 300)
    peel_box = (1010, 55, 690, 300)
    draw_axes(svg, tensile_box, (0, 11), (0, 30), [0, 2, 4, 6, 8, 10], [0, 10, 20, 30],
              "Displacement (mm)", "Tensile adhesion strength (kPa)")
    draw_curve(svg, tensile_fr, tensile_box, (0, 11), (0, 30), FR)
    draw_curve(svg, tensile_ptpu, tensile_box, (0, 11), (0, 30), PTPU)
    draw_legend(svg, 565, 82)

    draw_axes(svg, peel_box, (0, 80), (0, 400), [0, 20, 40, 60, 80], [0, 100, 200, 300, 400],
              "Displacement (mm)", "Peel strength, 2F/W (N/m)")
    draw_curve(svg, peel_fr, peel_box, (0, 80), (0, 400), FR)
    draw_curve(svg, peel_ptpu, peel_box, (0, 80), (0, 400), PTPU)
    draw_legend(svg, 1480, 82)

    svg.append(text(25, 475, "d", size=36, anchor="start", weight=700))
    svg.append(text(165, 488, "Strain = 50%", size=22, anchor="start", weight=600))
    draw_inset(svg, cycle_raw, 500, (455, 452, 330, 135))
    draw_inset(svg, cycle_raw, 1500, (1015, 452, 330, 135))

    main_box = (150, 625, 1500, 205)
    xlim, ylim = (0, 2000), (1.4, 2.0)
    draw_axes(svg, main_box, xlim, ylim, [0, 500, 1000, 1500, 2000], [1.4, 1.6, 1.8, 2.0],
              "Number of cycles", "Resistance (Ω)", tick_size=18)
    upper_points = [map_point(x, y, main_box, xlim, ylim) for x, y in upper]
    lower_points = [map_point(x, y, main_box, xlim, ylim) for x, y in lower]
    polygon = upper_points + list(reversed(lower_points))
    svg.append(f'<path d="{path_d(polygon)} Z" fill="{PALE}" stroke="none" opacity="0.55"/>')
    svg.append(f'<path d="{path_d(upper_points)}" fill="none" stroke="{FR}" stroke-width="3.7" stroke-linejoin="round" stroke-linecap="round"/>')
    svg.append(f'<path d="{path_d(lower_points)}" fill="none" stroke="{LOWER}" stroke-width="3.7" stroke-linejoin="round" stroke-linecap="round"/>')
    svg.append(line(565, 650, 620, 650, stroke=FR, width=5.0))
    svg.append(text(637, 657, "Upper envelope", size=20, anchor="start", weight=500))
    svg.append(line(935, 650, 990, 650, stroke=LOWER, width=5.0))
    svg.append(text(1007, 657, "Lower envelope", size=20, anchor="start", weight=500))

    svg.append('</svg>')
    OUTPUT.write_text("\n".join(svg), encoding="utf-8")
    print(OUTPUT)


if __name__ == "__main__":
    main()
