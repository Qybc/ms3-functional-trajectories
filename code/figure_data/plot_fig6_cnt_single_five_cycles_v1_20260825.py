#!/usr/bin/env python3
"""Draw five complete cycles for the supplied CNT/CB single-layer record."""

from __future__ import annotations

from html import escape
from pathlib import Path

import numpy as np
from openpyxl import load_workbook
from scipy.signal import find_peaks


HERE = Path(__file__).resolve().parent
PACKAGE_ROOT = HERE.parent.parent
RAW = PACKAGE_ROOT / "data/fig4_conductive_network/raw/CNT_CB_single_layer_raw.xlsx"
OUT = HERE / "outputs/fig4_cnt_cb_single_five_cycles.svg"

BLACK = "#000000"
GRID = "#e4e7e9"
COLORS = {"30%": "#2777b8", "60%": "#d64045", "100%": "#8a3fb0"}


def read_five_cycles(sheet, x_col, start, horizon, minimum_trough_distance_s):
    rows = []
    for x, y in sheet.iter_rows(
        min_row=3, min_col=x_col, max_col=x_col + 1, values_only=True
    ):
        if isinstance(x, (int, float)) and isinstance(y, (int, float)):
            if start <= float(x) <= start + horizon:
                rows.append((float(x) - start, float(y)))
    values = np.asarray(rows, dtype=float)
    if values.size == 0:
        raise ValueError(f"No numeric window in columns {x_col}:{x_col + 1}")

    r0 = float(np.quantile(values[:, 1], 0.01))
    response = np.maximum((values[:, 1] - r0) / r0 * 100.0, 0.0)
    sample_interval = float(np.median(np.diff(values[:, 0])))
    troughs, _ = find_peaks(
        -response,
        distance=max(1, int(minimum_trough_distance_s / sample_interval)),
        prominence=max(2.0, float(np.quantile(response, 0.90)) * 0.12),
    )
    usable = troughs[values[troughs, 0] >= 0.5]
    if len(usable) < 6:
        raise ValueError(f"Could not identify five complete cycles in columns {x_col}:{x_col + 1}")
    first, last = int(usable[0]), int(usable[5])
    measured_time = values[first : last + 1, 0] - values[first, 0]
    cycle_axis = measured_time / measured_time[-1] * 5.0
    response = response[first : last + 1].copy()
    response[0] = 0.0
    response[-1] = 0.0
    return cycle_axis, response


def path_d(x, y):
    return " ".join(
        [f"M{x[0]:.2f},{y[0]:.2f}"]
        + [f"L{xx:.2f},{yy:.2f}" for xx, yy in zip(x[1:], y[1:])]
    )


def text(x, y, value, cls, anchor="middle"):
    return (
        f'<text x="{x:.2f}" y="{y:.2f}" class="{cls}" '
        f'text-anchor="{anchor}">{escape(value)}</text>'
    )


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    workbook = load_workbook(RAW, read_only=True, data_only=True)
    sheet = workbook.active
    specifications = [
        ("30%", 1, 1000.0, 520.0, 45.0),
        ("60%", 3, 1000.0, 360.0, 28.0),
        ("100%", 5, 800.0, 140.0, 8.0),
    ]
    traces = []
    for index, (label, column, start, horizon, distance) in enumerate(specifications):
        x, y = read_five_cycles(sheet, column, start, horizon, distance)
        traces.append((label, x + 5.0 * index, y))
    workbook.close()

    width, height = 2200, 1020
    left, top, plot_width, plot_height = 235, 125, 1790, 690
    xlim, ylim = (0.0, 15.0), (0.0, 700.0)

    def mx(values):
        return left + (np.asarray(values) - xlim[0]) / (xlim[1] - xlim[0]) * plot_width

    def my(values):
        return top + plot_height - (np.asarray(values) - ylim[0]) / (ylim[1] - ylim[0]) * plot_height

    parts = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<style>',
        "text{font-family:'Avenir Next',Avenir,Arial,sans-serif;fill:#000000}",
        '.tick{font-size:38px;font-weight:500}',
        '.axis{font-size:48px;font-weight:650}',
        '.legend{font-size:42px;font-weight:650}',
        '</style>',
        f'<rect x="0" y="0" width="{width}" height="{height}" fill="white"/>',
        f'<clipPath id="plotclip"><rect x="{left}" y="{top}" width="{plot_width}" height="{plot_height}"/></clipPath>',
    ]

    for value in (200, 400, 600):
        yy = float(my([value])[0])
        parts.append(
            f'<line x1="{left}" y1="{yy:.2f}" x2="{left + plot_width}" y2="{yy:.2f}" '
            f'stroke="{GRID}" stroke-width="3"/>'
        )
    for label, x, y in traces:
        px, py = mx(x), my(y)
        parts.append(
            f'<path d="{path_d(px, py)}" fill="none" stroke="{COLORS[label]}" '
            f'stroke-width="7.5" stroke-linejoin="round" stroke-linecap="round" '
            f'clip-path="url(#plotclip)"/>'
        )
    parts.append(
        f'<rect x="{left}" y="{top}" width="{plot_width}" height="{plot_height}" '
        f'fill="none" stroke="{BLACK}" stroke-width="8"/>'
    )

    for value in (0, 5, 10, 15):
        xx = float(mx([value])[0])
        parts.append(
            f'<line x1="{xx:.2f}" y1="{top + plot_height}" x2="{xx:.2f}" '
            f'y2="{top + plot_height + 16}" stroke="{BLACK}" stroke-width="7"/>'
        )
        parts.append(text(xx, top + plot_height + 62, str(value), "tick"))
    for value in (0, 200, 400, 600, 700):
        yy = float(my([value])[0])
        parts.append(
            f'<line x1="{left - 16}" y1="{yy:.2f}" x2="{left}" y2="{yy:.2f}" '
            f'stroke="{BLACK}" stroke-width="7"/>'
        )
        parts.append(text(left - 30, yy + 13, str(value), "tick", "end"))

    parts.append(text(left + plot_width / 2, 935, "Cycle sequence", "axis"))
    parts.append(
        f'<text x="62" y="{top + plot_height / 2}" class="axis" text-anchor="middle" '
        f'transform="rotate(-90 62 {top + plot_height / 2})">ΔR/R0 (%)</text>'
    )
    legend_x = [700, 1110, 1535]
    for xx, (label, _, _) in zip(legend_x, traces):
        parts.append(
            f'<line x1="{xx - 95}" y1="66" x2="{xx - 25}" y2="66" '
            f'stroke="{COLORS[label]}" stroke-width="8"/>'
        )
        parts.append(text(xx, 80, f"{label} strain", "legend", "start"))

    parts.append('</svg>')
    OUT.write_text("".join(parts), encoding="utf-8")
    print(OUT)


if __name__ == "__main__":
    main()
