#!/usr/bin/env python3
"""Render a full-width single-layer/dual-layer five-cycle comparison."""

from __future__ import annotations

import importlib.util
from html import escape
from pathlib import Path

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    assert spec.loader is not None
    spec.loader.exec_module(module)
    return module


SINGLE = load_module("fig6_single", HERE / "plot_fig6_cnt_single_five_cycles_v1_20260825.py")
DUAL = load_module("fig6_dual", HERE / "plot_fig6_dual_sequential_drr_v1_20260825.py")
OUT = HERE / "outputs/fig4_single_dual_fullwidth.svg"

BLACK = "#000000"
GRID = "#e4e7e9"
COLORS = {"30%": "#2777b8", "60%": "#d64045", "100%": "#8a3fb0"}


def text(x, y, value, cls, anchor="middle", rotate=None):
    transform = f' transform="rotate({rotate} {x} {y})"' if rotate is not None else ""
    return (
        f'<text x="{x:.2f}" y="{y:.2f}" class="{cls}" '
        f'text-anchor="{anchor}"{transform}>{escape(value)}</text>'
    )


def path_d(x, y):
    return " ".join(
        [f"M{x[0]:.2f},{y[0]:.2f}"]
        + [f"L{xx:.2f},{yy:.2f}" for xx, yy in zip(x[1:], y[1:])]
    )


def plot(parts, box, title, traces, ymax, yticks, clip_id):
    left, top, width, height = box

    def mx(value):
        return left + value / 15.0 * width

    def my(value):
        return top + height - value / ymax * height

    parts.append(text(left + width / 2, top - 48, title, "plot-title"))
    for value in yticks:
        if value not in (0, ymax):
            yy = my(value)
            parts.append(
                f'<line x1="{left}" y1="{yy:.2f}" x2="{left + width}" y2="{yy:.2f}" '
                f'stroke="{GRID}" stroke-width="3"/>'
            )
    parts.append(
        f'<clipPath id="{clip_id}"><rect x="{left}" y="{top}" width="{width}" height="{height}"/></clipPath>'
    )
    for label, x, y in traces:
        px = [mx(float(v)) for v in x]
        py = [my(float(v)) for v in y]
        parts.append(
            f'<path d="{path_d(px, py)}" fill="none" stroke="{COLORS[label]}" '
            f'stroke-width="8" stroke-linejoin="round" stroke-linecap="round" '
            f'clip-path="url(#{clip_id})"/>'
        )
    parts.append(
        f'<rect x="{left}" y="{top}" width="{width}" height="{height}" '
        f'fill="none" stroke="{BLACK}" stroke-width="9"/>'
    )
    for value in (0, 5, 10, 15):
        xx = mx(value)
        parts.append(
            f'<line x1="{xx:.2f}" y1="{top + height}" x2="{xx:.2f}" '
            f'y2="{top + height + 16}" stroke="{BLACK}" stroke-width="7"/>'
        )
        parts.append(text(xx, top + height + 58, str(value), "tick"))
    for value in yticks:
        yy = my(value)
        parts.append(
            f'<line x1="{left - 16}" y1="{yy:.2f}" x2="{left}" y2="{yy:.2f}" '
            f'stroke="{BLACK}" stroke-width="7"/>'
        )
        parts.append(text(left - 28, yy + 13, str(value), "tick", "end"))
    parts.append(text(left + width / 2, top + height + 112, "Cycle sequence", "axis"))
    parts.append(text(left - 115, top + height / 2, "ΔR/R0 (%)", "axis", rotate=-90))


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    single_workbook = load_workbook(SINGLE.RAW, read_only=True, data_only=True)
    single_sheet = single_workbook.active
    single_specs = [
        ("30%", 1, 1000.0, 520.0, 45.0),
        ("60%", 3, 1000.0, 360.0, 28.0),
        ("100%", 5, 800.0, 140.0, 8.0),
    ]
    single_traces = []
    for index, (label, column, start, horizon, distance) in enumerate(single_specs):
        x, y = SINGLE.read_five_cycles(single_sheet, column, start, horizon, distance)
        single_traces.append((label, x + 5.0 * index, y))
    single_workbook.close()

    dual_workbook = load_workbook(DUAL.RAW, read_only=True, data_only=True)
    dual_sheet = dual_workbook.active
    dual_specs = [
        ("30%", 1, 260.0, 3.0),
        ("60%", 3, 600.0, 6.0),
        ("100%", 5, 900.0, 8.0),
    ]
    dual_traces = []
    for index, (label, column, start, distance) in enumerate(dual_specs):
        x, y = DUAL.read_five_cycles(dual_sheet, column, start, distance)
        dual_traces.append((label, x + 5.0 * index, y))
    dual_workbook.close()

    width, height = 3000, 950
    parts = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{width}" height="{height}" viewBox="0 0 {width} {height}">',
        '<style>',
        "text{font-family:'Avenir Next',Avenir,Arial,sans-serif;fill:#000000}",
        '.tick{font-size:39px;font-weight:500}',
        '.axis{font-size:49px;font-weight:650}',
        '.plot-title{font-size:48px;font-weight:650}',
        '.legend{font-size:43px;font-weight:650}',
        '</style>',
        f'<rect x="0" y="0" width="{width}" height="{height}" fill="white"/>',
    ]

    legend_centres = [1120, 1500, 1900]
    for centre, label in zip(legend_centres, ("30%", "60%", "100%")):
        parts.append(
            f'<line x1="{centre - 105}" y1="55" x2="{centre - 35}" y2="55" '
            f'stroke="{COLORS[label]}" stroke-width="9"/>'
        )
        parts.append(text(centre, 70, f"{label} strain", "legend", "start"))

    plot(parts, (190, 180, 1190, 560), "CNT/CB single layer",
         single_traces, 700.0, (0, 200, 400, 600, 700), "clip1")
    plot(parts, (1640, 180, 1190, 560), "CNT/CB–Ag dual layer",
         dual_traces, 1500.0, (0, 500, 1000, 1500), "clip2")
    parts.append('</svg>')
    OUT.write_text("".join(parts), encoding="utf-8")
    print(OUT)


if __name__ == "__main__":
    main()
