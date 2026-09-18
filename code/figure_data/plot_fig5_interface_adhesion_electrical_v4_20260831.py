#!/usr/bin/env python3
"""Create a compact, editable Fig. 5 c-d material-characterization block.

Panel c contains three raw adhesion responses (lap shear, normal tensile
adhesion and 180-degree peel). Panel d combines the extension-to-electrical
failure response with the 2,000 nominal-cycle resistance test. All curves are
read directly from the expert-provided Excel workbooks. The SVG deliberately
contains no clip paths so the Illustrator file remains easy to edit.
"""

from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import load_workbook


HERE = Path(__file__).resolve().parent
PACKAGE_ROOT = HERE.parent.parent
ROOT = HERE / "outputs"
RAW = PACKAGE_ROOT / "data/fig5_cardiac_interface/raw"
OUTPUT = ROOT / "fig5_selected_interface_characterization_v7_fullrange_20260831.svg"

W, H = 1800, 1240
INK = "#111111"
TEAL = "#315D70"
ORANGE = "#E86F51"
PALE = "#DDE9ED"
COLORS = {
    "Fr-PA-5": "#A7B4BC",
    "Fr-PA-10": "#315D70",
    "Fr-PA-20": "#E66F51",
    "PTPU-10": "#F2A65A",
}


def numeric_pairs(path, x_col, y_col, start_row=1):
    ws = load_workbook(path, read_only=True, data_only=True).active
    values = []
    for row in ws.iter_rows(min_row=start_row, values_only=True):
        x = row[x_col] if x_col < len(row) else None
        y = row[y_col] if y_col < len(row) else None
        if isinstance(x, (int, float)) and isinstance(y, (int, float)):
            values.append((float(x), float(y)))
    return values


def text(x, y, value, size=20, anchor="middle", weight=400, rotate=None, fill=INK):
    transform = f' transform="rotate({rotate} {x:.2f} {y:.2f})"' if rotate is not None else ""
    return (
        f'<text x="{x:.2f}" y="{y:.2f}" font-size="{size}" font-weight="{weight}" '
        f'text-anchor="{anchor}" fill="{fill}"{transform}>{escape(str(value))}</text>'
    )


def line(x1, y1, x2, y2, stroke=INK, width=3.0, dash=None):
    dash_attr = f' stroke-dasharray="{dash}"' if dash else ""
    return (
        f'<line x1="{x1:.2f}" y1="{y1:.2f}" x2="{x2:.2f}" y2="{y2:.2f}" '
        f'stroke="{stroke}" stroke-width="{width}"{dash_attr}/>'
    )


def rect(x, y, w, h, fill="white", stroke=INK, width=3.0):
    return (
        f'<rect x="{x:.2f}" y="{y:.2f}" width="{w:.2f}" height="{h:.2f}" '
        f'fill="{fill}" stroke="{stroke}" stroke-width="{width}"/>'
    )


def map_point(x, y, box, xlim, ylim):
    bx, by, bw, bh = box
    px = bx + (x - xlim[0]) / (xlim[1] - xlim[0]) * bw
    py = by + bh - (y - ylim[0]) / (ylim[1] - ylim[0]) * bh
    return px, py


def clamp(value, limits):
    return max(limits[0], min(limits[1], value))


def path_d(points, close=False):
    if not points:
        return ""
    d = "M " + " L ".join(f"{x:.2f},{y:.2f}" for x, y in points)
    return d + (" Z" if close else "")


def curve(svg, values, box, xlim, ylim, color, width=3.2, clamp_to_frame=False):
    pts = []
    for x, y in values:
        if x < xlim[0] or x > xlim[1]:
            continue
        if clamp_to_frame:
            y = clamp(y, ylim)
        elif y < ylim[0] or y > ylim[1]:
            continue
        pts.append(map_point(x, y, box, xlim, ylim))
    svg.append(
        f'<path d="{path_d(pts)}" fill="none" stroke="{color}" stroke-width="{width}" '
        f'stroke-linejoin="round" stroke-linecap="round"/>'
    )


def axes(svg, box, xlim, ylim, xticks, yticks, xlabel, ylabel, full_frame=True,
         tick_size=9, tick_font=20, label_font=26, grid=True):
    x, y, w, h = box
    if grid:
        for value in yticks[1:-1]:
            _, py = map_point(xlim[0], value, box, xlim, ylim)
            svg.append(line(x, py, x + w, py, stroke="#222222", width=1.7, dash="9 9"))
    if full_frame:
        svg.append(rect(x, y, w, h, fill="none", stroke=INK, width=4.4))
    else:
        svg.append(line(x, y, x, y + h, width=4.4))
        svg.append(line(x, y + h, x + w, y + h, width=4.4))
    for value in xticks:
        px, _ = map_point(value, ylim[0], box, xlim, ylim)
        svg.append(line(px, y + h, px, y + h + tick_size, width=3.4))
        svg.append(text(px, y + h + 30, f"{value:g}", size=tick_font))
    for value in yticks:
        _, py = map_point(xlim[0], value, box, xlim, ylim)
        svg.append(line(x - tick_size, py, x, py, width=3.4))
        svg.append(text(x - 14, py + 6, f"{value:g}", size=tick_font, anchor="end"))
    svg.append(text(x + w / 2, y + h + 65, xlabel, size=label_font, weight=600))
    svg.append(text(x - 63, y + h / 2, ylabel, size=label_font, weight=600, rotate=-90))


def adhesion_data():
    folder = RAW
    lap_path = folder / "lap_shear_raw.xlsx"
    tensile_path = folder / "tensile_adhesion_raw.xlsx"
    peel_path = folder / "peel_180_degree_raw.xlsx"
    # Assignments reproduce the legends of the expert-provided reference PNGs.
    lap = {
        "Fr-PA-20": numeric_pairs(lap_path, 0, 1),
        "Fr-PA-10": numeric_pairs(lap_path, 3, 4),
        "Fr-PA-5": numeric_pairs(lap_path, 7, 8),
        "PTPU-10": numeric_pairs(lap_path, 11, 12),
    }
    tensile = {
        "Fr-PA-10": numeric_pairs(tensile_path, 0, 1),
        "PTPU-10": numeric_pairs(tensile_path, 3, 4),
        "Fr-PA-20": numeric_pairs(tensile_path, 6, 7),
        "Fr-PA-5": numeric_pairs(tensile_path, 9, 10),
    }
    peel = {
        "Fr-PA-5": numeric_pairs(peel_path, 0, 1),
        "Fr-PA-10": numeric_pairs(peel_path, 3, 4),
        "Fr-PA-20": numeric_pairs(peel_path, 6, 7),
        "PTPU-10": numeric_pairs(peel_path, 9, 10),
    }
    return lap, tensile, peel


def draw_shared_legend(svg, y):
    names = ["Fr-PA-5", "Fr-PA-10", "Fr-PA-20", "PTPU-10"]
    starts = [410, 670, 945, 1225]
    for name, x in zip(names, starts):
        svg.append(line(x, y, x + 50, y, stroke=COLORS[name], width=7.0))
        svg.append(text(x + 64, y + 7, name, size=22, anchor="start"))


def draw_peel_legend(svg):
    x, y, w, h = 1535, 277, 187, 165
    svg.append(rect(x, y, w, h, fill="white", stroke="#777777", width=1.5))
    entries = ["Fr-PA-5", "Fr-PA-10", "Fr-PA-20", "PTPU-10"]
    for index, name in enumerate(entries):
        yy = y + 33 + index * 36
        svg.append(line(x + 17, yy, x + 61, yy, stroke=COLORS[name], width=8.0))
        svg.append(text(x + 73, yy + 8, name, size=21, anchor="start", weight=500))


def draw_adhesion(svg):
    lap, tensile, peel = adhesion_data()
    svg.append(text(42, 58, "c", size=36, anchor="start", weight=700))
    boxes = [(105, 115, 455, 360), (700, 115, 455, 360), (1295, 115, 455, 360)]
    specs = [
        (lap, boxes[0], (0, 4.5), (0, 50), [0, 1, 2, 3, 4], [0, 10, 20, 30, 40, 50],
         "Lap-shear strength (kPa)"),
        (tensile, boxes[1], (0, 11), (0, 30), [0, 2, 4, 6, 8, 10], [0, 10, 20, 30],
         "Tensile adhesion strength (kPa)"),
        (peel, boxes[2], (0, 80), (0, 400), [0, 20, 40, 60, 80], [0, 100, 200, 300, 400],
         "Peel strength, 2F/W (N/m)"),
    ]
    for curves, box, xlim, ylim, xticks, yticks, ylabel in specs:
        axes(svg, box, xlim, ylim, xticks, yticks, "Displacement (mm)", ylabel)
        for name in ["Fr-PA-5", "Fr-PA-10", "Fr-PA-20", "PTPU-10"]:
            curve(svg, curves[name], box, xlim, ylim, COLORS[name], width=4.5)
    draw_peel_legend(svg)


def percentile(values, fraction):
    ordered = sorted(values)
    position = fraction * (len(ordered) - 1)
    lower = int(position)
    upper = min(lower + 1, len(ordered) - 1)
    weight = position - lower
    return ordered[lower] * (1 - weight) + ordered[upper] * weight


def windowed_envelope(raw, total=2000, window_width=10):
    last_time = raw[-1][0]
    windows = [[] for _ in range(total // window_width)]
    for time_value, resistance in raw:
        nominal = time_value / last_time * total
        index = int(nominal // window_width)
        if 0 <= index < len(windows):
            windows[index].append(resistance)
    upper, lower = [], []
    for index, values in enumerate(windows):
        if len(values) < 20:
            continue
        x = index * window_width + window_width / 2
        upper.append((x, percentile(values, 0.95)))
        lower.append((x, percentile(values, 0.05)))
    return upper, lower


def draw_inset(svg, raw, center, box):
    last_time = raw[-1][0]
    selected = []
    for time_value, resistance in raw:
        nominal = time_value / last_time * 2000
        if center - 5 <= nominal <= center + 5:
            selected.append((nominal, resistance))
    xlim, ylim = (center - 5, center + 5), (1.4, 2.0)
    axes(svg, box, xlim, ylim, [center - 5, center, center + 5], [1.4, 1.7, 2.0], "", "",
         tick_size=6, tick_font=16, label_font=18)
    curve(svg, selected, box, xlim, ylim, TEAL, width=3.6, clamp_to_frame=True)


def draw_electrical(svg):
    svg.append(text(42, 690, "d", size=36, anchor="start", weight=700))
    folder = RAW
    extension = numeric_pairs(folder / "extension_resistance_raw.xlsx", 0, 1, 2)
    cyclic = numeric_pairs(folder / "cyclic_resistance_2000cycles_50pct_raw.xlsx", 0, 1, 2)

    left = (105, 760, 610, 350)
    axes(svg, left, (0, 1100), (0, 50), [0, 200, 400, 600, 800, 1000], [0, 10, 20, 30, 40, 50],
         "Strain (%)", "Resistance (Ω)")
    curve(svg, extension, left, (0, 1100), (0, 50), TEAL, width=4.8)
    svg.append(text(128, 792, "Extension rate = 50 mm/min", size=19, anchor="start", weight=500))

    svg.append(text(920, 716, "Strain = 50%", size=22, anchor="start", weight=600))
    draw_inset(svg, cyclic, 500, (920, 735, 300, 155))
    draw_inset(svg, cyclic, 1500, (1390, 735, 300, 155))

    main = (850, 925, 850, 185)
    axes(svg, main, (0, 2000), (1.4, 2.0), [0, 500, 1000, 1500, 2000], [1.4, 1.6, 1.8, 2.0],
         "Number of cycles", "Resistance (Ω)", tick_font=18, label_font=24)
    upper, lower = windowed_envelope(cyclic)
    upper_pts = [map_point(x, y, main, (0, 2000), (1.4, 2.0)) for x, y in upper]
    lower_pts = [map_point(x, y, main, (0, 2000), (1.4, 2.0)) for x, y in lower]
    polygon = upper_pts + list(reversed(lower_pts))
    svg.append(f'<path d="{path_d(polygon, close=True)}" fill="{PALE}" stroke="none" opacity="0.62"/>')
    svg.append(f'<path d="{path_d(upper_pts)}" fill="none" stroke="{TEAL}" stroke-width="3.8" stroke-linejoin="round"/>')
    svg.append(f'<path d="{path_d(lower_pts)}" fill="none" stroke="{ORANGE}" stroke-width="3.8" stroke-linejoin="round"/>')
    svg.append(line(1095, 953, 1138, 953, stroke=TEAL, width=4.0))
    svg.append(text(1150, 960, "Upper envelope", size=18, anchor="start"))
    svg.append(line(1390, 953, 1433, 953, stroke=ORANGE, width=4.0))
    svg.append(text(1445, 960, "Lower envelope", size=18, anchor="start"))


def main():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    if OUTPUT.exists():
        raise SystemExit(f"Refusing to overwrite existing output: {OUTPUT}")
    svg = [
        '<?xml version="1.0" encoding="UTF-8"?>',
        f'<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}" viewBox="0 0 {W} {H}">',
        '<defs><style><![CDATA[',
        "text { font-family: 'Avenir Next', Avenir, Arial, sans-serif; }",
        ']]></style></defs>',
        f'<rect x="0" y="0" width="{W}" height="{H}" fill="white"/>',
    ]
    draw_adhesion(svg)
    draw_electrical(svg)
    svg.append('</svg>')
    OUTPUT.write_text("\n".join(svg), encoding="utf-8")
    print(OUTPUT)


if __name__ == "__main__":
    main()
