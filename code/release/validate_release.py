#!/usr/bin/env python3
"""Validate integrity and key structural invariants of an MS3 release."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import subprocess
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook


REQUIRED = (
    "README.md",
    "VERSION.json",
    "CHANGELOG.md",
    "LICENSE",
    "LICENSE-DATA.md",
    ".zenodo.json",
    "THIRD_PARTY_NOTICES.md",
    "requirements-agent.txt",
    "source_data/MS3_Source_Data_v1_20260918.xlsx",
    "supplementary_data/MS3_Supplementary_Data_1_benchmark_audit_v1_20260918.xlsx",
    "data/fig2_agent_evaluation/benchmark_summary.csv",
    "data/fig2_agent_evaluation/fs_evidenceqa_per_item_scores.csv",
    "data/fig3_innovation_backtesting/nested_metrics.csv",
    "data/fig3_innovation_backtesting/strict_primary_trajectory_units.csv",
    "data/fig3_innovation_backtesting/new_complete_trajectories_2022_2025_row_level.csv",
    "data/fig3_innovation_backtesting/ranked_top100_candidates.csv",
    "schema/ms3_record.schema.json",
    "examples/ms3_record_example.json",
    "examples/ms3_record_real_selected_fields.json",
    "examples/ms3_record_real_coherent_case.json",
    "examples/ms3_record_real_partial_case.json",
    "examples/fs_evidenceqa_real_output_cases.json",
    "configs/agent_protocol_v1.json",
    "configs/evaluation_registry_v1.json",
    "manifests/RUN_PROVENANCE.json",
    "docs/INSTALL_AND_DEMO.md",
    "docs/PROVENANCE_AND_RIGHTS.md",
    "code/analysis/reproduce_core_results.py",
    "code/demo/build_demo_db.py",
    "code/release/build_source_data.py",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def worksheet_records(ws, header_row: int = 1) -> list[tuple]:
    headers = [cell.value for cell in ws[header_row] if cell.value is not None]
    records = []
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        values = tuple(row[: len(headers)])
        if any(value is not None for value in values):
            records.append(values)
    return records


def workbook_values(path: Path) -> dict[str, list[tuple]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    return {
        name: list(workbook[name].iter_rows(values_only=True))
        for name in workbook.sheetnames
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("release_root", nargs="?", type=Path, default=Path("."))
    parser.add_argument("--skip-manifest", action="store_true")
    parser.add_argument(
        "--deep",
        action="store_true",
        help="rebuild both workbooks from released CSV/JSON and compare every cell",
    )
    args = parser.parse_args()
    root = args.release_root.resolve()

    failures: list[str] = []

    forbidden_artifact_names = {".DS_Store"}
    for path in root.rglob("*"):
        if path.name in forbidden_artifact_names or path.name.startswith("._"):
            failures.append(f"forbidden filesystem artifact: {path.relative_to(root)}")
        if path.is_dir() and path.name == "__pycache__":
            failures.append(f"Python cache directory in release: {path.relative_to(root)}")
        if path.is_file() and path.suffix == ".pyc":
            failures.append(f"compiled Python cache in release: {path.relative_to(root)}")

    text_suffixes = {".md", ".py", ".json", ".yml", ".yaml", ".txt", ".cff"}
    sensitive_patterns = {
        "local macOS user path": re.compile("/" + "Users/" + r"[^/\s]+/"),
        "MS3 server path": re.compile("/mnt/" + "diskshare/"),
        "private-network address": re.compile(
            r"\b(?:100\.87\.207\." + r"13|172\.18\.124\." + r"(?:251|252))\b"
        ),
    }
    for path in root.rglob("*"):
        if not path.is_file() or path.suffix.lower() not in text_suffixes:
            continue
        content = path.read_text(encoding="utf-8", errors="ignore")
        for label, pattern in sensitive_patterns.items():
            if pattern.search(content):
                failures.append(
                    f"{label} found in public text: {path.relative_to(root)}"
                )
    for relative in REQUIRED:
        if not (root / relative).is_file():
            failures.append(f"missing required file: {relative}")

    benchmark = read_csv(root / "data/fig2_agent_evaluation/benchmark_summary.csv")
    if len(benchmark) != 24:
        failures.append(f"benchmark summary has {len(benchmark)} rows; expected 24")
    observed = {(row["benchmark"], row["method"]) for row in benchmark}
    if len(observed) != 24:
        failures.append("benchmark and method pairs are not unique")

    fs_items = read_csv(root / "data/fig2_agent_evaluation/fs_evidenceqa_per_item_scores.csv")
    if len(fs_items) != 720:
        failures.append(f"FS EvidenceQA has {len(fs_items)} rows; expected 720")
    case_ids = {row["case_id"] for row in fs_items}
    methods = {row["method"] for row in fs_items}
    if len(case_ids) != 120 or len(methods) != 6:
        failures.append(
            f"FS EvidenceQA dimensions are {len(case_ids)} cases x {len(methods)} methods; expected 120 x 6"
        )

    backtests = read_csv(root / "data/fig3_innovation_backtesting/nested_metrics.csv")
    if len(backtests) != 24:
        failures.append(f"historical backtest has {len(backtests)} rows; expected 24")

    strict_units = read_csv(
        root
        / "data/fig3_innovation_backtesting/strict_primary_trajectory_units.csv"
    )
    if len(strict_units) != 4127:
        failures.append(
            f"strict trajectory table has {len(strict_units)} rows; expected 4127"
        )

    recent_paths = read_csv(
        root
        / "data/fig3_innovation_backtesting/new_complete_trajectories_2022_2025_row_level.csv"
    )
    if len(recent_paths) != 1302:
        failures.append(
            f"row-level 2022-2025 trajectory table has {len(recent_paths)} rows; expected 1302"
        )
    if len({row["trajectory"] for row in recent_paths}) != len(recent_paths):
        failures.append("row-level 2022-2025 trajectories are not unique")
    expected_modes = {
        "Higher-order recombination": 1021,
        "Material implementation": 140,
        "Transduction novelty": 78,
        "Application transfer": 28,
        "Multi-layer novelty": 35,
    }
    observed_modes: dict[str, int] = {}
    for row in recent_paths:
        observed_modes[row["mode"]] = observed_modes.get(row["mode"], 0) + 1
    if observed_modes != expected_modes:
        failures.append(
            f"row-level innovation modes differ from expected values: {observed_modes}"
        )

    ranked_top100 = read_csv(
        root / "data/fig3_innovation_backtesting/ranked_top100_candidates.csv"
    )
    if len(ranked_top100) != 1200:
        failures.append(
            f"ranked Top-100 candidate table has {len(ranked_top100)} rows; expected 1200"
        )
    ranking_groups: dict[tuple[str, str], list[int]] = {}
    for row in ranked_top100:
        ranking_groups.setdefault((row["cutoff"], row["method"]), []).append(
            int(row["rank"])
        )
    if len(ranking_groups) != 12 or any(
        sorted(ranks) != list(range(1, 101)) for ranks in ranking_groups.values()
    ):
        failures.append("ranked Top-100 candidate groups are incomplete or non-contiguous")

    workbook = load_workbook(
        root / "source_data/MS3_Source_Data_v1_20260918.xlsx", read_only=True, data_only=True
    )
    required_sheets = {
        "README",
        "Fig1_corpus",
        "Fig2_benchmarks",
        "Fig3_backtests",
        "Fig4_d_traces",
        "Fig5_e_ECG",
    }
    absent_sheets = required_sheets - set(workbook.sheetnames)
    if absent_sheets:
        failures.append(f"Source Data workbook missing sheets: {sorted(absent_sheets)}")

    supplementary_workbook = load_workbook(
        root
        / "supplementary_data/MS3_Supplementary_Data_1_benchmark_audit_v1_20260918.xlsx",
        read_only=True,
        data_only=True,
    )
    required_supplementary_sheets = {
        "README",
        "Benchmark_summary",
        "FS_method_summary",
        "FS_paired_bootstrap",
        "Matched_rep_ablation",
        "FS_per_item",
        "Historical_top100",
        "Ranked_top100",
        "Trajectory_1302",
        "Example_single",
        "Example_cross",
    }
    absent_supplementary_sheets = required_supplementary_sheets - set(
        supplementary_workbook.sheetnames
    )
    if absent_supplementary_sheets:
        failures.append(
            "Supplementary Data workbook missing sheets: "
            f"{sorted(absent_supplementary_sheets)}"
        )
    if supplementary_workbook["FS_per_item"].max_row != 724:
        failures.append(
            "Supplementary Data FS_per_item sheet does not contain 720 data rows"
        )
    if supplementary_workbook["Trajectory_1302"].max_row != 1306:
        failures.append(
            "Supplementary Data Trajectory_1302 sheet does not contain 1302 data rows"
        )
    if supplementary_workbook["Ranked_top100"].max_row != 1204:
        failures.append(
            "Supplementary Data Ranked_top100 sheet does not contain 1200 data rows"
        )

    source_benchmarks = worksheet_records(workbook["Fig2_benchmarks"])
    supplementary_benchmarks = worksheet_records(
        supplementary_workbook["Benchmark_summary"], header_row=4
    )
    if source_benchmarks != supplementary_benchmarks:
        failures.append("benchmark rows differ between Source Data and Supplementary Data")

    source_fs_items = worksheet_records(workbook["Fig2_FS_per_item"])
    supplementary_fs_items = worksheet_records(
        supplementary_workbook["FS_per_item"], header_row=4
    )
    if source_fs_items != supplementary_fs_items:
        failures.append("FS EvidenceQA per-item rows differ between the two workbooks")

    source_backtests = worksheet_records(workbook["Fig3_backtests"])
    supplementary_backtests = worksheet_records(
        supplementary_workbook["Historical_top100"], header_row=4
    )
    if source_backtests != supplementary_backtests:
        failures.append("historical Top-100 rows differ between the two workbooks")

    example = json.loads((root / "examples/ms3_record_example.json").read_text(encoding="utf-8"))
    evidence_ids = {row["evidence_id"] for row in example["evidence"]}
    for mechanism in example["mechanisms"]:
        unknown = set(mechanism["evidence_chain"]) - evidence_ids
        if unknown:
            failures.append(f"example record has dangling evidence identifiers: {sorted(unknown)}")

    protocol = json.loads(
        (root / "configs/agent_protocol_v1.json").read_text(encoding="utf-8")
    )
    if protocol.get("max_tool_calls") != 6 or len(protocol.get("tools", [])) != 3:
        failures.append("agent protocol must define three tools and a six-call limit")

    registry = json.loads(
        (root / "configs/evaluation_registry_v1.json").read_text(encoding="utf-8")
    )
    registered_sizes = {
        row["name"]: row["n_items"] for row in registry.get("evaluations", [])
    }
    expected_sizes = {
        "LitQA2-FullText": 75,
        "ScholarQA-CS2": 100,
        "SciCUEval Materials": 200,
        "FS EvidenceQA": 120,
    }
    if registered_sizes != expected_sizes:
        failures.append(
            f"evaluation registry sizes differ from expected values: {registered_sizes}"
        )

    provenance = json.loads(
        (root / "manifests/RUN_PROVENANCE.json").read_text(encoding="utf-8")
    )
    for artifact in provenance.get("artifacts", []):
        relative = artifact.get("relative_path")
        expected = artifact.get("sha256")
        if not relative or not expected:
            failures.append("run provenance contains an incomplete artifact entry")
            continue
        target = root / relative
        if not target.is_file():
            failures.append(f"run provenance target missing: {relative}")
        elif sha256(target) != expected:
            failures.append(f"run provenance checksum mismatch: {relative}")

    manifest_path = root / "manifests/MANIFEST.tsv"
    if manifest_path.is_file() and not args.skip_manifest:
        with manifest_path.open(encoding="utf-8", newline="") as handle:
            for row in csv.DictReader(handle, delimiter="\t"):
                relative = row.get("path") or row.get("relative_path")
                expected = row.get("sha256")
                if not relative or not expected:
                    continue
                target = root / relative
                if not target.is_file():
                    failures.append(f"manifest target missing: {relative}")
                elif sha256(target) != expected:
                    failures.append(f"manifest checksum mismatch: {relative}")

    if args.deep:
        with tempfile.TemporaryDirectory(prefix="ms3-release-check-") as directory:
            temporary = Path(directory)
            rebuilt_source = temporary / "Source_Data.xlsx"
            rebuilt_supplementary = temporary / "Supplementary_Data_1.xlsx"
            commands = (
                (root / "code/release/build_source_data.py", rebuilt_source),
                (
                    root / "code/release/build_supplementary_data.py",
                    rebuilt_supplementary,
                ),
            )
            for script, output in commands:
                subprocess.run(
                    [
                        sys.executable,
                        str(script),
                        "--release-root",
                        str(root),
                        "--output",
                        str(output),
                    ],
                    check=True,
                    capture_output=True,
                    text=True,
                )
            if workbook_values(rebuilt_source) != workbook_values(
                root / "source_data/MS3_Source_Data_v1_20260918.xlsx"
            ):
                failures.append("deep check: rebuilt Source Data differs from release workbook")
            if workbook_values(rebuilt_supplementary) != workbook_values(
                root
                / "supplementary_data/MS3_Supplementary_Data_1_benchmark_audit_v1_20260918.xlsx"
            ):
                failures.append(
                    "deep check: rebuilt Supplementary Data differs from release workbook"
                )

    if failures:
        print("validation: FAIL")
        for failure in failures:
            print(f"- {failure}")
        raise SystemExit(1)
    print("validation: PASS")
    print(f"benchmark rows: {len(benchmark)}")
    print(f"FS EvidenceQA: {len(case_ids)} cases x {len(methods)} methods")
    print(f"historical backtest rows: {len(backtests)}")
    print(f"strict paper-trajectory units: {len(strict_units)}")
    print(f"new complete trajectories, 2022-2025: {len(recent_paths)}")
    print(f"ranked historical Top-100 candidates: {len(ranked_top100)}")
    print(f"Source Data sheets: {len(workbook.sheetnames)}")
    print(f"Supplementary Data sheets: {len(supplementary_workbook.sheetnames)}")
    if args.deep:
        print("deep workbook rebuild checks: PASS")


if __name__ == "__main__":
    main()
