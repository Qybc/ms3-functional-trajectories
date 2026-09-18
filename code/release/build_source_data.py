#!/usr/bin/env python3
"""Rebuild the Nature-style Source Data workbook from released CSV files."""

from __future__ import annotations

import argparse
import csv
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


SHEETS = (
    ("Fig1_corpus", "data/fig1_framework/corpus_scale.csv"),
    ("Fig1_connectivity", "data/fig1_framework/trajectory_connectivity.csv"),
    ("Fig1_grounding", "data/fig1_framework/evidence_grounding.csv"),
    ("Fig1_extraction", "data/fig1_framework/extraction_evaluation.csv"),
    ("Fig2_benchmarks", "data/fig2_agent_evaluation/benchmark_summary.csv"),
    ("Fig2_FS_summary", "data/fig2_agent_evaluation/fs_evidenceqa_method_summary.csv"),
    ("Fig2_FS_subgroups", "data/fig2_agent_evaluation/fs_evidenceqa_subgroups.csv"),
    ("Fig2_FS_dimensions", "data/fig2_agent_evaluation/fs_evidenceqa_dimension_subgroups.csv"),
    ("Fig2_FS_paired", "data/fig2_agent_evaluation/fs_evidenceqa_paired_bootstrap.csv"),
    ("Fig2_ablation", "data/fig2_agent_evaluation/matched_role_graph_ablation.csv"),
    ("Fig2_FS_per_item", "data/fig2_agent_evaluation/fs_evidenceqa_per_item_scores.csv"),
    ("Fig3_annual", "data/fig3_innovation_backtesting/annual_composition.csv"),
    ("Fig3_new_paths", "data/fig3_innovation_backtesting/new_complete_trajectories_2022_2025.csv"),
    ("Fig3_reobservation", "data/fig3_innovation_backtesting/four_year_reobservation.csv"),
    ("Fig3_backtests", "data/fig3_innovation_backtesting/top100_closure_summary.csv"),
    ("Fig4_d_traces", "data/fig4_conductive_network/representative_five_cycle_traces.csv"),
    ("Fig5_c_tensile", "data/fig5_cardiac_interface/tensile_adhesion_selected_groups.csv"),
    ("Fig5_c_peel", "data/fig5_cardiac_interface/peel_selected_groups.csv"),
    ("Fig5_d_extension", "data/fig5_cardiac_interface/extension_resistance.csv"),
    ("Fig5_d_cycles", "data/fig5_cardiac_interface/cyclic_resistance_envelope.csv"),
    ("Fig5_e_ECG", "data/fig5_cardiac_interface/porcine_ecg_display_points.csv"),
    ("Fig5_e_HR", "data/fig5_cardiac_interface/porcine_heart_rate_display_points.csv"),
)

INTEGER = re.compile(r"^-?(?:0|[1-9][0-9]*)$")
FLOAT = re.compile(r"^-?(?:(?:0|[1-9][0-9]*)\.[0-9]+|[0-9]+(?:\.[0-9]+)?[eE][+-]?[0-9]+)$")


def typed(value: str):
    if value == "":
        return None
    if value == "True":
        return True
    if value == "False":
        return False
    if INTEGER.fullmatch(value):
        return int(value)
    if FLOAT.fullmatch(value):
        return float(value)
    return value


def add_csv_sheet(workbook: Workbook, root: Path, name: str, relative: str) -> None:
    path = root / relative
    with path.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        rows = list(reader)
    if not rows:
        raise ValueError(f"empty CSV: {relative}")

    ws = workbook.create_sheet(title=name)
    for index, row in enumerate(rows, start=1):
        ws.append(row if index == 1 else [typed(value) for value in row])

    fill = PatternFill("solid", fgColor="D9EAF2")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in range(1, ws.max_column + 1):
        observed = [len(str(ws.cell(row=row, column=column).value or ""))
                    for row in range(1, min(ws.max_row, 501) + 1)]
        ws.column_dimensions[get_column_letter(column)].width = min(max(observed) + 2, 42)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--release-root", type=Path,
                        default=Path(__file__).resolve().parents[2])
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    root = args.release_root.resolve()

    workbook = Workbook()
    ws = workbook.active
    ws.title = "README"
    readme = (
        ("Dataset", "Source data for Evidence-linked functional trajectories for literature-driven scientific innovation"),
        ("Version", "v1, 18 September 2026"),
        ("Scope", "Numeric data underlying plotted results in main Figs. 1-5"),
        ("FS EvidenceQA", "Per-item numeric judgments are included; full answers and copyrighted evidence packets are not redistributed"),
        ("Public benchmarks", "Complete-set aggregate scores are included; item-level model outputs are not included in the shared release"),
        ("Images", "Raw SEM/EDS images are distributed as separate files under data/fig4_conductive_network/raw and data/fig5_cardiac_interface/raw"),
        ("Note", "Representative experimental traces are descriptive; independent-replicate counts are reported in the manuscript"),
    )
    for key, value in readme:
        ws.append((key, value))
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    for name, relative in SHEETS:
        add_csv_sheet(workbook, root, name, relative)

    output = args.output.resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    print(f"wrote {len(workbook.sheetnames)} sheets to {output}")


if __name__ == "__main__":
    main()
