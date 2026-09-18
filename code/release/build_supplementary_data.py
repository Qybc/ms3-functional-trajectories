#!/usr/bin/env python3
"""Build the reviewer-facing MS3 Supplementary Data workbook.

The workbook collects frozen numeric evaluation outputs that are too large for
the Supplementary Information PDF. It intentionally excludes copyrighted
article text, model credentials and unavailable final per-item archives for the
three public benchmarks.
"""

from __future__ import annotations

import argparse
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


BLUE = "DDEBF7"
DARK_BLUE = "1F4E78"
LIGHT_BLUE = "EAF3F8"
GREY = "F2F2F2"
WHITE = "FFFFFF"
THIN_GREY = Side(style="thin", color="D9E2F3")


def _read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _coerce(value: str):
    if value is None:
        return None
    text = str(value).strip()
    if text == "":
        return ""
    if text.lower() == "true":
        return True
    if text.lower() == "false":
        return False
    try:
        if any(marker in text.lower() for marker in (".", "e")):
            return float(text)
        return int(text)
    except ValueError:
        return text


def _title(ws, title: str, subtitle: str | None = None) -> int:
    ws.cell(1, 1, title)
    ws.cell(1, 1).font = Font(size=15, bold=True, color=DARK_BLUE)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    row = 2
    if subtitle:
        ws.cell(row, 1, subtitle)
        ws.cell(row, 1).font = Font(size=10, italic=True, color="666666")
        ws.cell(row, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        ws.row_dimensions[row].height = 32
        row += 1
    return row + 1


def _style_table(ws, header_row: int, last_row: int, last_col: int) -> None:
    for cell in ws[header_row]:
        if cell.column <= last_col:
            cell.fill = PatternFill("solid", fgColor=DARK_BLUE)
            cell.font = Font(bold=True, color=WHITE)
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            cell.border = Border(bottom=THIN_GREY)
    for row in ws.iter_rows(
        min_row=header_row + 1, max_row=last_row, min_col=1, max_col=last_col
    ):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN_GREY)
    ws.freeze_panes = f"A{header_row + 1}"
    ws.auto_filter.ref = (
        f"A{header_row}:{get_column_letter(last_col)}{last_row}"
    )


def _write_records(
    ws,
    title: str,
    subtitle: str,
    records: list[dict],
    widths: dict[str, float] | None = None,
) -> None:
    header_row = _title(ws, title, subtitle)
    if not records:
        ws.cell(header_row, 1, "No records")
        return
    headers = list(records[0].keys())
    for col, name in enumerate(headers, 1):
        ws.cell(header_row, col, name)
    for row_idx, record in enumerate(records, header_row + 1):
        for col_idx, name in enumerate(headers, 1):
            ws.cell(row_idx, col_idx, _coerce(record.get(name)))
    _style_table(ws, header_row, header_row + len(records), len(headers))
    widths = widths or {}
    for col_idx, name in enumerate(headers, 1):
        width = widths.get(name)
        if width is None:
            sample = [len(str(record.get(name, ""))) for record in records[:100]]
            width = min(max([len(name), *sample]) + 2, 36)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.sheet_view.showGridLines = False


def _example_records(path: Path) -> list[dict]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    mapping = payload["mapping"]
    records = []
    for evaluation in payload["judgment"]["evaluations"]:
        label = evaluation["label"]
        records.append(
            {
                "case_id": payload["case_id"],
                "blind_label": label,
                "method_after_unblinding": mapping[label],
                "scientific_correctness_0_4": evaluation["scientific_correctness"],
                "source_groundedness_0_4": evaluation["source_groundedness"],
                "evidence_completeness_0_4": evaluation["evidence_completeness"],
                "boundary_control_0_4": evaluation["boundary_control"],
                "unsupported_claim_count": evaluation["unsupported_claim_count"],
                "most_important_error": evaluation["most_important_error"],
                "justification": evaluation["justification"],
                "audit_packet_adequate": payload["judgment"][
                    "audit_packet_adequate"
                ],
            }
        )
    return records


def _add_readme(wb: Workbook) -> None:
    ws = wb.active
    ws.title = "README"
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 102
    ws["A1"] = "MS3 Supplementary Data 1"
    ws["A1"].font = Font(size=17, bold=True, color=DARK_BLUE)
    ws.merge_cells("A1:B1")
    rows = [
        (
            "Manuscript",
            "Evidence-linked functional trajectories for literature-driven scientific innovation",
        ),
        ("Version", "v1, 18 September 2026"),
        (
            "Purpose",
            "Reviewer-facing numeric audit tables for the reasoning evaluation and historical backtesting. Main-figure plotting values remain in the separate Source Data workbook.",
        ),
        (
            "Scope",
            "Includes complete aggregate results for six displayed systems on four evaluations; 120-question FS EvidenceQA per-item scores; subgroup, dimension and paired-bootstrap results; the matched representation ablation; two frozen blinded scoring examples; all four Top-100 historical backtests; and the 1,302 row-level trajectories underlying the innovation-composition result.",
        ),
        (
            "Public-benchmark boundary",
            "LitQA2 (75), ScholarQA-CS2 (100) and SciCUEval Materials (200) are reported as complete-set aggregate evaluations without redistributing their benchmark content. FS EvidenceQA includes the released per-item numeric judgments used for its analysis.",
        ),
        (
            "Copyright boundary",
            "No publisher PDF, copyrighted figure, long article excerpt or complete proprietary audit packet is redistributed. Case sheets contain numeric judgments and evaluator explanations only.",
        ),
        (
            "FS EvidenceQA composite",
            "Unweighted mean of scientific correctness, source groundedness, evidence completeness and evidence-boundary control, each scored 0-4. normalized_percent = composite/4 x 100. Unsupported substantive claims are counted separately.",
        ),
        (
            "Paired bootstrap",
            "10,000 question-level paired resamples. All method scores for a question remain paired within each resample. Intervals are percentile 95% confidence intervals and secondary contrasts are not multiplicity-adjusted.",
        ),
        (
            "Matched ablation",
            "Role Graph versus MS3 is a separate matched representation experiment. The frozen blind archive contained seven systems; the paper and per-item sheet display six, excluding the archived LLaMat run.",
        ),
        (
            "Files",
            "The machine-readable CSV/JSON sources and evaluation rubric are distributed in data/fig2_agent_evaluation and code/agent of the MS3 data-and-code package.",
        ),
    ]
    start = 3
    for idx, (key, value) in enumerate(rows, start):
        ws.cell(idx, 1, key)
        ws.cell(idx, 2, value)
        ws.cell(idx, 1).font = Font(bold=True, color=DARK_BLUE)
        ws.cell(idx, 1).fill = PatternFill("solid", fgColor=BLUE)
        ws.cell(idx, 1).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(idx, 2).alignment = Alignment(wrap_text=True, vertical="top")
        ws.cell(idx, 1).border = Border(bottom=THIN_GREY)
        ws.cell(idx, 2).border = Border(bottom=THIN_GREY)
        ws.row_dimensions[idx].height = 50 if len(value) > 180 else 36

    index_start = start + len(rows) + 2
    ws.cell(index_start, 1, "Sheet")
    ws.cell(index_start, 2, "Contents")
    sheet_rows = [
        ("Benchmark_summary", "Six-system aggregate scores for all four evaluations."),
        ("FS_method_summary", "Overall FS EvidenceQA dimensions, composite, unsupported claims and rank."),
        ("FS_subgroups", "Prespecified single-paper and cross-paper summaries."),
        ("FS_dimensions", "Dimension-level subgroup results."),
        ("FS_paired_bootstrap", "Paired MS3 contrasts with 95% confidence intervals."),
        ("Matched_rep_ablation", "Matched Role Graph versus MS3 comparison."),
        ("FS_per_item", "720 numeric rows: 120 cases x 6 displayed systems."),
        ("Example_single", "Frozen blinded scoring example nv3_single_01_04."),
        ("Example_cross", "Frozen blinded scoring example nv3_cross_06_07."),
        ("Historical_top100", "Four cutoffs, two closure endpoints and three relation-informed rankings."),
        ("Ranked_top100", "Candidate-level Top-100 lists for each cutoff and relation-informed ranking."),
        ("Trajectory_1302", "Row-level normalized trajectories underlying the 2022-2025 innovation-composition result."),
    ]
    for col in (1, 2):
        cell = ws.cell(index_start, col)
        cell.fill = PatternFill("solid", fgColor=DARK_BLUE)
        cell.font = Font(bold=True, color=WHITE)
    for row_idx, (sheet, description) in enumerate(sheet_rows, index_start + 1):
        ws.cell(row_idx, 1, sheet)
        ws.cell(row_idx, 2, description)
        ws.cell(row_idx, 1).font = Font(bold=True)
        ws.cell(row_idx, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        ws.cell(row_idx, 1).border = Border(bottom=THIN_GREY)
        ws.cell(row_idx, 2).border = Border(bottom=THIN_GREY)
        ws.cell(row_idx, 2).alignment = Alignment(wrap_text=True)
    ws.freeze_panes = "A3"


def build(release_root: Path, output: Path) -> None:
    data_root = release_root / "data"
    agent_root = data_root / "fig2_agent_evaluation"
    history_root = data_root / "fig3_innovation_backtesting"

    wb = Workbook()
    _add_readme(wb)

    sheets = [
        (
            "Benchmark_summary",
            "Complete aggregate benchmark matrix",
            "Scores are percentages. Missing outputs were scored as incorrect in the frozen summaries.",
            agent_root / "benchmark_summary.csv",
        ),
        (
            "FS_method_summary",
            "FS EvidenceQA overall method summary",
            "All 120 questions; six displayed systems.",
            agent_root / "fs_evidenceqa_method_summary.csv",
        ),
        (
            "FS_subgroups",
            "FS EvidenceQA prespecified subgroups",
            "Single-paper and cross-paper questions were defined before analysis (60 each).",
            agent_root / "fs_evidenceqa_subgroups.csv",
        ),
        (
            "FS_dimensions",
            "FS EvidenceQA dimension-level subgroup results",
            "Each dimension uses the 0-4 rubric recorded in code/agent/FS_EVIDENCEQA_REPORTING_RUBRIC.md.",
            agent_root / "fs_evidenceqa_dimension_subgroups.csv",
        ),
        (
            "FS_paired_bootstrap",
            "FS EvidenceQA paired bootstrap contrasts",
            "MS3 minus comparator; 10,000 paired question-level resamples.",
            agent_root / "fs_evidenceqa_paired_bootstrap.csv",
        ),
        (
            "Matched_rep_ablation",
            "Matched representation ablation",
            "Role Graph and MS3 used the same controller, corpus and six-call budget. This is separate from the six-system blind evaluation.",
            agent_root / "matched_role_graph_ablation.csv",
        ),
        (
            "FS_per_item",
            "FS EvidenceQA per-item numeric judgments",
            "One row per question and displayed method (120 x 6 = 720). Article excerpts and complete answers are excluded from the public workbook.",
            agent_root / "fs_evidenceqa_per_item_scores.csv",
        ),
        (
            "Historical_top100",
            "Historical backtesting: Top-100 outcomes",
            "Structure-plus-evidence is retrospective because its formulation was refined after inspection of historical outcomes.",
            history_root / "top100_closure_summary.csv",
        ),
        (
            "Ranked_top100",
            "Historical backtesting: candidate-level Top-100 rankings",
            "Four cutoffs x three relation-informed methods x 100 candidates. Outcome columns reproduce the Top-100 closure counts in the summary sheet; structure-plus-evidence is retrospective.",
            history_root / "ranked_top100_candidates.csv",
        ),
        (
            "Trajectory_1302",
            "New complete trajectories first observed in 2022-2025",
            "The 1,302 unique normalized paths underlying the reported 1,021/1,302 (78.4%) recombination result; no publisher full text or long excerpts are included.",
            history_root / "new_complete_trajectories_2022_2025_row_level.csv",
        ),
    ]
    for name, title, subtitle, path in sheets:
        ws = wb.create_sheet(name)
        _write_records(ws, title, subtitle, _read_csv(path))

    annotated = agent_root / "annotated_examples"
    for name, filename, title in (
        (
            "Example_single",
            "nv3_single_01_04_blind_judgment.json",
            "Blinded scoring example: single-paper case nv3_single_01_04",
        ),
        (
            "Example_cross",
            "nv3_cross_06_07_blind_judgment.json",
            "Blinded scoring example: cross-paper case nv3_cross_06_07",
        ),
    ):
        ws = wb.create_sheet(name)
        _write_records(
            ws,
            title,
            "Method identities were restored only after scoring; evaluator explanations are reproduced from the frozen judgment file.",
            _example_records(annotated / filename),
            {"most_important_error": 58, "justification": 110},
        )
        for row in range(5, ws.max_row + 1):
            ws.row_dimensions[row].height = 95

    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.page_margins.left = 0.3
        ws.page_margins.right = 0.3
        ws.page_margins.top = 0.5
        ws.page_margins.bottom = 0.5

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--release-root",
        type=Path,
        default=Path(__file__).resolve().parents[2],
    )
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.release_root.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
