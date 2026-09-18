#!/usr/bin/env python3
"""Build the MS3 source-data release from frozen local artifacts.

This script does not modify manuscript source files or original experiment
directories.  It writes derived CSV files and the Nature-style Source Data
workbook inside the independent release directory that contains this script.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import importlib.util
import json
import math
import os
import random
import shutil
from pathlib import Path
from statistics import mean

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


METHODS = {
    "ms3_agent_deepseek": "MS3 Agent",
    "deepseek_v4_pro_evidence": "DeepSeek V4 Pro",
    "kimi_k3_evidence": "Kimi K3",
    "glm_5_2_evidence": "GLM-5.2",
    "paperqa2_deepseek": "PaperQA2",
    "openscholar_style_deepseek": "OpenScholar",
}

PUBLIC_BENCHMARKS = [
    ("LitQA2-FullText", "exact-label accuracy", 75,
     [44.0, 32.0, 33.3, 16.0, 4.0, 17.3]),
    ("ScholarQA-CS2", "official global average", 100,
     [66.0, 31.9, 49.5, 43.5, 31.1, 54.0]),
    ("SciCUEval Materials", "eight-stratum macro-average", 200,
     [98.0, 91.0, 84.0, 91.5, 40.5, 48.5]),
    ("FS EvidenceQA", "normalized 0-4 composite", 120,
     [82.6, 66.4, 77.2, 72.0, 67.9, 55.2]),
]

DISPLAY_METHOD_ORDER = [
    "MS3 Agent", "DeepSeek V4 Pro", "Kimi K3", "GLM-5.2",
    "PaperQA2", "OpenScholar",
]


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_csv(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    if not rows:
        raise ValueError(f"No rows for {path}")
    fields: list[str] = []
    for row in rows:
        for key in row:
            if key not in fields:
                fields.append(key)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(rows)


def copy_file(source: Path, target: Path) -> None:
    if not source.exists():
        raise FileNotFoundError(source)
    target.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, target)


def load_module(name: str, path: Path):
    spec = importlib.util.spec_from_file_location(name, path)
    if spec is None or spec.loader is None:
        raise RuntimeError(f"Cannot import {path}")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def percentile(values: list[float], fraction: float) -> float:
    ordered = sorted(values)
    position = fraction * (len(ordered) - 1)
    lo = int(position)
    hi = min(lo + 1, len(ordered) - 1)
    weight = position - lo
    return ordered[lo] * (1 - weight) + ordered[hi] * weight


def bootstrap_paired(differences: list[float], replicates: int = 10_000,
                     seed: int = 42) -> tuple[float, float]:
    rng = random.Random(seed)
    n = len(differences)
    draws = []
    for _ in range(replicates):
        draws.append(mean(differences[rng.randrange(n)] for _ in range(n)))
    return percentile(draws, 0.025), percentile(draws, 0.975)


def numeric_pairs(path: Path, x_col: int, y_col: int,
                  min_row: int = 2) -> list[tuple[float, float]]:
    ws = load_workbook(path, read_only=True, data_only=True).active
    pairs = []
    for row in ws.iter_rows(min_row=min_row, values_only=True):
        x = row[x_col - 1] if x_col <= len(row) else None
        y = row[y_col - 1] if y_col <= len(row) else None
        if isinstance(x, (int, float)) and isinstance(y, (int, float)):
            pairs.append((float(x), float(y)))
    return pairs


def build_fig1(data_root: Path) -> dict[str, list[dict]]:
    corpus = [
        {"measure": "papers", "value": 13689},
        {"measure": "evidence_items", "value": 131083},
        {"measure": "trajectory_records", "value": 26648},
        {"measure": "complete_trajectories", "value": 15703},
    ]
    connectivity = [
        {"trajectory_class": "Complete", "percentage": 58.93},
        {"trajectory_class": "Contiguous partial", "percentage": 30.33},
        {"trajectory_class": "Evidence-gapped", "percentage": 10.74},
    ]
    grounding = [
        {"grounding_profile": "Text only", "percentage": 26.1},
        {"grounding_profile": "Figure/caption", "percentage": 11.5},
        {"grounding_profile": "Table/quantitative", "percentage": 18.7},
        {"grounding_profile": "Text + both", "percentage": 43.7},
    ]
    extraction = []
    values = {
        "Direct": (69.2, 70.8),
        "Single-stage SFT": (78.0, 77.2),
        "E/T/M": (91.7, 95.6),
    }
    for method, (completeness, quality) in values.items():
        extraction.extend([
            {"method": method, "metric": "Record completeness", "score_percent": completeness},
            {"method": method, "metric": "Effective record quality", "score_percent": quality},
        ])
    out = data_root / "fig1_framework"
    write_csv(out / "corpus_scale.csv", corpus)
    write_csv(out / "trajectory_connectivity.csv", connectivity)
    write_csv(out / "evidence_grounding.csv", grounding)
    write_csv(out / "extraction_evaluation.csv", extraction)
    return {"Fig1_corpus": corpus, "Fig1_connectivity": connectivity,
            "Fig1_grounding": grounding, "Fig1_extraction": extraction}


def build_fig2(repo: Path, data_root: Path) -> dict[str, list[dict]]:
    out = data_root / "fig2_agent_evaluation"
    benchmark_rows = []
    for benchmark, metric, n, values in PUBLIC_BENCHMARKS:
        for method, score in zip(DISPLAY_METHOD_ORDER, values):
            benchmark_rows.append({
                "benchmark": benchmark,
                "metric": metric,
                "n_items": n,
                "method": method,
                "score_percent": score,
                "missing_outputs_scored_incorrect": True,
            })
    write_csv(out / "benchmark_summary.csv", benchmark_rows)

    blind_dir = (repo / "latest/experiments/covered120_7methods_20260816/artifacts"
                 / "blind_qwen30b_seven_methods_covered120_v3")
    judgments = json.loads((blind_dir / "blind_answer_judgments.json").read_text())
    summary = json.loads((blind_dir / "blind_answer_summary.json").read_text())
    blinding_metadata = {
        key: judgments[key]
        for key in (
            "status", "judge_api_base", "judge_model",
            "judge_is_same_provider_as_answer_backbone",
            "method_identity_hidden", "methods",
        )
        if key in judgments
    }
    (out / "blinding_metadata.json").write_text(
        json.dumps(blinding_metadata, indent=2, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    for case_id in ("nv3_single_01_04", "nv3_cross_06_07"):
        copy_file(blind_dir / "case_cache" / f"{case_id}.json",
                  out / "annotated_examples" / f"{case_id}_blind_judgment.json")
    per_item = []
    scores_by_method: dict[str, dict[str, float]] = {key: {} for key in METHODS}
    for record in judgments["records"]:
        case_id = record["case_id"]
        for method_id, display in METHODS.items():
            evaluation = record["method_evaluations"][method_id]
            dimensions = [
                evaluation["scientific_correctness"],
                evaluation["source_groundedness"],
                evaluation["evidence_completeness"],
                evaluation["boundary_control"],
            ]
            composite = sum(dimensions) / 4.0
            scores_by_method[method_id][case_id] = composite
            per_item.append({
                "case_id": case_id,
                "task_type": record["task_type"],
                "audit_packet_adequate": record["audit_packet_adequate"],
                "method": display,
                "scientific_correctness_0_4": dimensions[0],
                "source_groundedness_0_4": dimensions[1],
                "evidence_completeness_0_4": dimensions[2],
                "boundary_control_0_4": dimensions[3],
                "composite_0_4": composite,
                "normalized_percent": composite / 4.0 * 100.0,
                "unsupported_claim_count": evaluation["unsupported_claim_count"],
            })
    per_item.sort(key=lambda row: (row["case_id"], DISPLAY_METHOD_ORDER.index(row["method"])))
    write_csv(out / "fs_evidenceqa_per_item_scores.csv", per_item)

    summary_rows = []
    for method_id, display in METHODS.items():
        item = summary["methods"][method_id]
        summary_rows.append({
            "method": display,
            "n_items": summary["case_count"],
            "scientific_correctness_0_4": item["scientific_correctness"],
            "source_groundedness_0_4": item["source_groundedness"],
            "evidence_completeness_0_4": item["evidence_completeness"],
            "boundary_control_0_4": item["boundary_control"],
            "composite_0_4": item["composite_0_to_4"],
            "normalized_percent": item["composite_0_to_4"] / 4.0 * 100.0,
            "unsupported_claims_per_question": item["unsupported_claim_count"],
            "first_place_count": item["first_place_count"],
            "mean_rank": item["mean_rank"],
        })
    summary_rows.sort(key=lambda row: DISPLAY_METHOD_ORDER.index(row["method"]))
    write_csv(out / "fs_evidenceqa_method_summary.csv", summary_rows)

    subgroup_rows = []
    for subgroup, methods in summary["by_task_type"].items():
        for method_id, display in METHODS.items():
            item = methods[method_id]
            method_items = [r for r in per_item if r["task_type"] == subgroup and r["method"] == display]
            subgroup_rows.append({
                "task_type": subgroup,
                "method": display,
                "n_items": len(method_items),
                "composite_0_4": item["composite_0_to_4"],
                "normalized_percent": item["composite_0_to_4"] / 4.0 * 100.0,
                "unsupported_claims_per_question": item["unsupported_claim_count"],
            })
    write_csv(out / "fs_evidenceqa_subgroups.csv", subgroup_rows)

    comparisons = []
    focal = scores_by_method["ms3_agent_deepseek"]
    for method_id, display in METHODS.items():
        if method_id == "ms3_agent_deepseek":
            continue
        diffs = [focal[case] - scores_by_method[method_id][case] for case in sorted(focal)]
        frozen = summary["paired_focal_differences"][
            f"ms3_agent_deepseek_minus_{method_id}"
        ]["composite_0_to_4"]
        lo, hi = frozen["ci95"]
        positive, ties, negative = frozen["positive_ties_negative"]
        comparisons.append({
            "contrast": f"MS3 Agent - {display}",
            "n_paired_items": len(diffs),
            "mean_difference_0_4": frozen["mean"],
            "bootstrap_95ci_low": lo,
            "bootstrap_95ci_high": hi,
            "bootstrap_replicates": 10000,
            "positive_items": positive,
            "tied_items": ties,
            "negative_items": negative,
            "source": "frozen blind-evaluation summary",
        })
    write_csv(out / "fs_evidenceqa_paired_bootstrap.csv", comparisons)

    dimension_rows = []
    for task_type in ("single_paper", "cross_paper"):
        for display in DISPLAY_METHOD_ORDER:
            rows = [r for r in per_item if r["task_type"] == task_type and r["method"] == display]
            for metric in ("scientific_correctness_0_4", "source_groundedness_0_4",
                           "evidence_completeness_0_4", "boundary_control_0_4"):
                dimension_rows.append({
                    "task_type": task_type,
                    "method": display,
                    "metric": metric,
                    "mean_0_4": mean(float(row[metric]) for row in rows),
                    "normalized_percent": mean(float(row[metric]) for row in rows) / 4.0 * 100.0,
                    "n_items": len(rows),
                })
    write_csv(out / "fs_evidenceqa_dimension_subgroups.csv", dimension_rows)

    ablation = [
        {"metric": "Final-workspace source recall", "unit": "percent", "Role Graph": 61.3, "MS3": 82.9},
        {"metric": "All claims evidence-grounded", "unit": "percent", "Role Graph": 69.2, "MS3": 90.0},
        {"metric": "Unsupported claims per question", "unit": "count", "Role Graph": 2.275, "MS3": 1.567},
        {"metric": "Tool calls per question", "unit": "count", "Role Graph": 4.79, "MS3": 3.88},
    ]
    for row in ablation:
        row["MS3_minus_Role_Graph"] = row["MS3"] - row["Role Graph"]
        row["evaluation_note"] = "Matched representation ablation; separate from the seven-method blind FS evaluation"
    write_csv(out / "matched_role_graph_ablation.csv", ablation)
    copy_file(blind_dir / "blind_answer_summary.json", out / "fs_evidenceqa_blind_summary.json")
    return {
        "Fig2_benchmarks": benchmark_rows,
        "Fig2_FS_summary": summary_rows,
        "Fig2_FS_subgroups": subgroup_rows,
        "Fig2_FS_dimensions": dimension_rows,
        "Fig2_FS_paired": comparisons,
        "Fig2_ablation": ablation,
        "Fig2_FS_per_item": per_item,
    }


def build_fig3(repo: Path, data_root: Path) -> dict[str, list[dict]]:
    out = data_root / "fig3_innovation_backtesting"
    annual_values = {
        2020: (33.1, 45.6), 2021: (42.2, 42.5), 2022: (47.7, 43.8),
        2023: (47.7, 42.8), 2024: (52.0, 42.8), 2025: (53.4, 42.6),
    }
    annual = []
    for year, (existing, recombination) in annual_values.items():
        annual.extend([
            {"year": year, "mode": "Existing trajectory", "share_percent": existing},
            {"year": year, "mode": "Complete-path recombination", "share_percent": recombination},
            {"year": year, "mode": "Local-link novelty", "share_percent": round(100 - existing - recombination, 1)},
        ])
    write_csv(out / "annual_composition.csv", annual)

    new_paths = [
        {"class": "All three local links pre-observed", "count": 1021, "percent": 78.4},
        {"class": "Material-to-Sensor only", "count": 140, "percent": 10.8},
        {"class": "Sensor-to-Signal only", "count": 78, "percent": 6.0},
        {"class": "Signal-to-System only", "count": 28, "percent": 2.2},
        {"class": "Multi-layer novelty", "count": 35, "percent": 2.7},
    ]
    write_csv(out / "new_complete_trajectories_2022_2025.csv", new_paths)

    reobservation = [
        {"mode": "Complete-path recombination", "eligible_n": 656, "four_year_reobserved_percent": 31.6},
        {"mode": "Material implementation", "eligible_n": 174, "four_year_reobserved_percent": 17.8},
        {"mode": "Application transfer", "eligible_n": 30, "four_year_reobserved_percent": 20.0},
        {"mode": "Transduction", "eligible_n": 55, "four_year_reobserved_percent": 3.6},
        {"mode": "Multi-layer novelty", "eligible_n": 57, "four_year_reobserved_percent": 7.0},
    ]
    write_csv(out / "four_year_reobservation.csv", reobservation)

    hist = (repo / "latest/experiments/fig4_nested_evidence_20260823"
            / "nested_temporal_evidence_rankfusion_v3_20260823/nested_evaluation")
    copy_file(hist / "nested_metrics.csv", out / "nested_metrics.csv")
    copy_file(hist / "popularity_matched_metrics.csv", out / "popularity_matched_metrics.csv")
    copy_file(hist / "popularity_matched_pairs.csv", out / "popularity_matched_pairs.csv")
    copy_file(hist / "evaluation_summary.json", out / "evaluation_summary.json")
    nested_rows = list(csv.DictReader((hist / "nested_metrics.csv").open(encoding="utf-8")))
    top100 = []
    for row in nested_rows:
        positives = int(row["positives"])
        universe = int(row["candidate_universe"])
        top100.append({
            "cutoff": int(row["cutoff"]),
            "endpoint": row["endpoint"],
            "model": row["model"],
            "candidate_universe": universe,
            "future_positives": positives,
            "top100_hits": int(row["hits_at_100"]),
            "top100_precision": float(row["precision_at_100"]),
            "random_expected_hits_at_100": positives / universe * 100.0,
            "average_precision": float(row["average_precision"]),
            "selected_evidence_composite": row["selected_evidence_composite"],
            "selected_fusion_mode": row["selected_fusion_mode"],
            "selected_weight": float(row["selected_weight"]),
        })
    write_csv(out / "top100_closure_summary.csv", top100)
    return {"Fig3_annual": annual, "Fig3_new_paths": new_paths,
            "Fig3_reobservation": reobservation, "Fig3_backtests": top100}


def build_fig4(repo: Path, data_root: Path) -> dict[str, list[dict]]:
    out = data_root / "fig4_conductive_network"
    raw_out = out / "raw"
    source = repo / "latest/source/figures/fig6_source_20260825"
    copy_file(source / "signal or sensor 电学信号/CNT single layer.xlsx",
              raw_out / "CNT_CB_single_layer_raw.xlsx")
    copy_file(source / "signal or sensor 电学信号/dual layer CNT+Ag.xlsx",
              raw_out / "CNT_CB_Ag_dual_layer_raw.xlsx")
    for path in sorted((source / "Materials 电镜").glob("*.tif")):
        safe_name = path.name.replace(" ", "_")
        copy_file(path, raw_out / "SEM" / safe_name)

    fig_dir = repo / "latest/source/figures"
    single = load_module("release_fig4_single", fig_dir / "plot_fig6_cnt_single_five_cycles_v1_20260825.py")
    dual = load_module("release_fig4_dual", fig_dir / "plot_fig6_dual_sequential_drr_v1_20260825.py")
    rows = []
    workbook = load_workbook(single.RAW, read_only=True, data_only=True)
    sheet = workbook.active
    for index, (label, column, start, horizon, distance) in enumerate([
        ("30%", 1, 1000.0, 520.0, 45.0),
        ("60%", 3, 1000.0, 360.0, 28.0),
        ("100%", 5, 800.0, 140.0, 8.0),
    ]):
        x, y = single.read_five_cycles(sheet, column, start, horizon, distance)
        rows.extend({"architecture": "CNT/CB single layer", "strain": label,
                     "cycle_sequence": float(xx + 5 * index), "delta_R_over_R0_percent": float(yy)}
                    for xx, yy in zip(x, y))
    workbook.close()
    workbook = load_workbook(dual.RAW, read_only=True, data_only=True)
    sheet = workbook.active
    for index, (label, column, start, distance) in enumerate([
        ("30%", 1, 260.0, 3.0), ("60%", 3, 600.0, 6.0), ("100%", 5, 900.0, 8.0),
    ]):
        x, y = dual.read_five_cycles(sheet, column, start, distance)
        rows.extend({"architecture": "CNT/CB-Ag dual layer", "strain": label,
                     "cycle_sequence": float(xx + 5 * index), "delta_R_over_R0_percent": float(yy)}
                    for xx, yy in zip(x, y))
    workbook.close()
    write_csv(out / "representative_five_cycle_traces.csv", rows)
    return {"Fig4_d_traces": rows}


def parse_point_string(value: str) -> list[tuple[float, float]]:
    points = []
    for token in value.split(";"):
        if not token:
            continue
        x, y = token.split(",")
        points.append((float(x), float(y)))
    return points


def build_fig5(repo: Path, data_root: Path, raw_root: Path) -> dict[str, list[dict]]:
    out = data_root / "fig5_cardiac_interface"
    raw_out = out / "raw"
    sensor = raw_root / "Sensor-传感材料数据"
    mappings = {
        sensor / "传感材料粘附数据/Lap-shear strength.xlsx": raw_out / "lap_shear_raw.xlsx",
        sensor / "传感材料粘附数据/Tensile strength.xlsx": raw_out / "tensile_adhesion_raw.xlsx",
        sensor / "传感材料粘附数据/粘附模型180peel 模型.xlsx": raw_out / "peel_180_degree_raw.xlsx",
        sensor / "传感材料电学数据/以50mmmin速率 拉断传感材料电阻变化.xlsx": raw_out / "extension_resistance_raw.xlsx",
        sensor / "传感材料电学数据/2000次 50%应变重复循环电学原始数据.xlsx": raw_out / "cyclic_resistance_2000cycles_50pct_raw.xlsx",
    }
    for source, target in mappings.items():
        copy_file(source, target)

    micro = raw_root / "Materials/纤维材料电镜照片/GA-IN-PA"
    for path in sorted(micro.glob("*.tiff")):
        safe = (path.name.replace(" ", "_").replace("α", "alpha").replace(",", "_"))
        copy_file(path, raw_out / "SEM_EDS" / safe)

    tensile_path = sensor / "传感材料粘附数据/Tensile strength.xlsx"
    peel_path = sensor / "传感材料粘附数据/粘附模型180peel 模型.xlsx"
    tensile = []
    for group, x_col, y_col in [("Fr-PA-10", 4, 5), ("PTPU-10", 10, 11)]:
        tensile.extend({"group": group, "displacement_mm": x, "tensile_adhesion_kPa": y}
                       for x, y in numeric_pairs(tensile_path, x_col, y_col))
    peel = []
    for group, x_col, y_col in [("Fr-PA-10", 4, 5), ("PTPU-10", 10, 11)]:
        peel.extend({"group": group, "displacement_mm": x, "peel_strength_2F_over_W_N_per_m": y}
                    for x, y in numeric_pairs(peel_path, x_col, y_col))
    extension = [{"strain_percent": x, "resistance_ohm": y} for x, y in numeric_pairs(
        sensor / "传感材料电学数据/以50mmmin速率 拉断传感材料电阻变化.xlsx", 1, 2)]
    cyclic_pairs = numeric_pairs(
        sensor / "传感材料电学数据/2000次 50%应变重复循环电学原始数据.xlsx", 1, 2)
    last_time = cyclic_pairs[-1][0]
    windows: list[list[float]] = [[] for _ in range(200)]
    for time_value, resistance in cyclic_pairs:
        coordinate = time_value / last_time * 2000
        index = int(coordinate // 10)
        if 0 <= index < len(windows):
            windows[index].append(resistance)
    cycle_envelope = []
    for index, values in enumerate(windows):
        if len(values) < 20:
            continue
        cycle_envelope.append({
            "cycle_window_midpoint": index * 10 + 5,
            "lower_5th_percentile_ohm": percentile(values, 0.05),
            "upper_95th_percentile_ohm": percentile(values, 0.95),
            "n_raw_points": len(values),
        })
    write_csv(out / "tensile_adhesion_selected_groups.csv", tensile)
    write_csv(out / "peel_selected_groups.csv", peel)
    write_csv(out / "extension_resistance.csv", extension)
    write_csv(out / "cyclic_resistance_envelope.csv", cycle_envelope)

    trace_path = repo / "latest/source/figures/fig5_panel_e_porcine_ecg_hr_trace_data_v8_20260907.txt"
    copy_file(trace_path, raw_out / "display_trace_payload_v8.txt")
    ecg_rows, hr_rows = [], []
    with trace_path.open(encoding="utf-8") as handle:
        for line in handle:
            if not line.strip() or line.startswith("#"):
                continue
            fields = line.rstrip("\n").split("|")
            label, duration, amp, plot_min, plot_max, source_name, block, start, stop, note, ecg, hr = fields
            for x, y in parse_point_string(ecg):
                ecg_rows.append({
                    "condition": label, "relative_time_fraction": x,
                    "relative_time_s": x * float(duration), "ecg_mV": y,
                    "source_file": source_name, "source_block": int(block),
                    "source_start_s": float(start), "source_stop_s": float(stop),
                    "display_note": note,
                })
            for x, y in parse_point_string(hr):
                hr_rows.append({
                    "condition": label, "relative_time_fraction": x,
                    "relative_time_s": x * float(duration), "heart_rate_bpm": y,
                    "smoothing": "centred 6-s moving median",
                })
    write_csv(out / "porcine_ecg_display_points.csv", ecg_rows)
    write_csv(out / "porcine_heart_rate_display_points.csv", hr_rows)
    return {
        "Fig5_c_tensile": tensile, "Fig5_c_peel": peel,
        "Fig5_d_extension": extension, "Fig5_d_cycles": cycle_envelope,
        "Fig5_e_ECG": ecg_rows, "Fig5_e_HR": hr_rows,
    }


def add_sheet(workbook: Workbook, name: str, rows: list[dict]) -> None:
    ws = workbook.create_sheet(title=name[:31])
    if not rows:
        ws.append(["No rows"])
        return
    headers = []
    for row in rows:
        for key in row:
            if key not in headers:
                headers.append(key)
    ws.append(headers)
    fill = PatternFill("solid", fgColor="D9EAF2")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row in rows:
        ws.append([row.get(key) for key in headers])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for index, header in enumerate(headers, 1):
        observed = [len(str(header))]
        for row in rows[:500]:
            value = row.get(header)
            if value is not None:
                observed.append(len(str(value)))
        ws.column_dimensions[get_column_letter(index)].width = min(max(observed) + 2, 42)


def build_workbook(path: Path, sheets: dict[str, list[dict]]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "README"
    readme = [
        ("Dataset", "Source data for Evidence-linked functional trajectories for literature-driven scientific innovation"),
        ("Version", "v1, 18 September 2026"),
        ("Scope", "Numeric data underlying plotted results in main Figs. 1-5"),
        ("FS EvidenceQA", "Per-item numeric judgments are included; full answers and copyrighted evidence packets are not redistributed"),
        ("Public benchmarks", "Frozen complete-set aggregate scores are included; public-benchmark per-item outputs are not part of the verified public snapshot"),
        ("Images", "Raw SEM/EDS images are distributed as separate files under data/fig4_conductive_network/raw and data/fig5_cardiac_interface/raw"),
        ("Note", "Representative experimental traces are descriptive; independent-replicate counts are reported in the manuscript"),
    ]
    for key, value in readme:
        ws.append([key, value])
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 110
    for row in ws.iter_rows():
        row[0].font = Font(bold=True)
        row[1].alignment = Alignment(wrap_text=True, vertical="top")
    for name, rows in sheets.items():
        add_sheet(wb, name, rows)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def build_manifest(package: Path) -> None:
    rows = []
    for path in sorted(package.rglob("*")):
        if not path.is_file() or path.name == "MANIFEST.tsv":
            continue
        rows.append({
            "relative_path": str(path.relative_to(package)),
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
        })
    manifest = package / "manifests/MANIFEST.tsv"
    manifest.parent.mkdir(parents=True, exist_ok=True)
    with manifest.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=["relative_path", "bytes", "sha256"], delimiter="\t")
        writer.writeheader()
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--repo-root", type=Path, default=Path(__file__).resolve().parents[5])
    parser.add_argument("--raw-root", type=Path, default=None)
    args = parser.parse_args()
    repo = args.repo_root.resolve()
    package = Path(__file__).resolve().parents[2]
    data_root = package / "data"

    sheets = {}
    sheets.update(build_fig1(data_root))
    sheets.update(build_fig2(repo, data_root))
    sheets.update(build_fig3(repo, data_root))
    sheets.update(build_fig4(repo, data_root))
    raw_root = args.raw_root or (Path(os.environ["MS3_RAW_ROOT"]) if "MS3_RAW_ROOT" in os.environ else None)
    if raw_root is None:
        parser.error("provide --raw-root or set MS3_RAW_ROOT")
    sheets.update(build_fig5(repo, data_root, raw_root.resolve()))
    build_workbook(package / "source_data/MS3_Source_Data_v1_20260918.xlsx", sheets)
    build_manifest(package)
    print(f"Built release at {package}")


if __name__ == "__main__":
    main()
